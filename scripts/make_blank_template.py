"""Create a privacy-clean quota template from the supplied original workbook."""

import argparse
import json
import os
from pathlib import Path
import re
import tempfile
from xml.dom import minidom
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.packaging.custom import CustomPropertyList


def sanitize_package_metadata(path):
    path = Path(path)
    with ZipFile(path) as source:
        original = source.read("xl/workbook.xml")
        document = minidom.parseString(original)
        removed = 0
        for node in list(document.getElementsByTagNameNS("*", "absPath")):
            node.parentNode.removeChild(node)
            removed += 1
        if not removed:
            return 0
        cleaned = document.toxml(encoding="utf-8")
        with tempfile.TemporaryDirectory(prefix="quota-template-metadata-", dir=path.parent) as temp:
            staged = Path(temp) / path.name
            with ZipFile(staged, "w") as destination:
                for entry in source.infolist():
                    destination.writestr(
                        entry, cleaned if entry.filename == "xl/workbook.xml"
                        else source.read(entry.filename),
                    )
            # Windows must close the source archive before replacing the file.
            source.close()
            os.replace(staged, path)
    return removed


def find_sheet(workbook, name):
    matches = [
        sheet for sheet in workbook
        if re.sub(r"\s+", "", sheet.title) == re.sub(r"\s+", "", name)
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one worksheet named {name!r}")
    return matches[0]


def make_template(source, output):
    source = Path(source).resolve()
    output = Path(output).resolve()
    if source == output:
        raise ValueError("The original workbook must not be overwritten")
    if output.exists():
        raise FileExistsError(output)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx source workbooks are supported")
    workbook = load_workbook(source, keep_links=False)
    requests = find_sheet(workbook, "요청목록")
    common = find_sheet(workbook, "공통정보")
    reference = find_sheet(workbook, "입력포맷")
    choices = find_sheet(workbook, "폼선택값")
    validation_count = len(requests.data_validations.dataValidation)
    defined_name_count = len(workbook.defined_names)
    required_headers = {
        1: "실행(Y/N)", 2: "요청ID", 6: "폼 모델명(정확한 표기)",
        8: "신청 총 TPM(기존+추가)", 10: "구독 ID(행별 지정 시)",
    }
    for column, expected in required_headers.items():
        if requests.cell(3, column).value != expected:
            raise ValueError(f"Unexpected request header in column {column}")
    source_values = {
        str(common.cell(row, 1).value): str(common.cell(row, 2).value)
        for row in range(2, common.max_row + 1)
        if common.cell(row, 2).value not in (None, "")
    }
    replacements = {
        "First Name": "ExampleFirst",
        "Last Name": "ExampleLast",
        "Company Email": "requester@example.com",
        "Company Name": "Example Corporation",
        "Company Address": "123 Example Street",
        "Company City": "Example City",
        "Company Postal Code": "00000",
        "Company Country": "Example Country",
        "Subscription Id": "00000000-0000-0000-0000-000000000000",
        "Justification": "Describe your workload and requested peak TPM.",
        "구독 이름(참고용)": "example-subscription",
    }
    for row in range(4, requests.max_row + 1):
        for column in (1, 2, 3, 4, 5, 6, 7, 8, 10, 11):
            cell = requests.cell(row, column)
            cell.value = None
            cell.hyperlink = None
    for row in range(2, common.max_row + 1):
        common.cell(row, 2).value = None
        common.cell(row, 2).hyperlink = None
    requests.cell(2, 1).value = requests.cell(2, 1).value.replace(
        "실행=Y인 행만 처리.",
        "실행=Y는 제출 대상 표시이며, 실제 전송 전 미리보기 승인이 필요합니다.",
    )
    for row in reference:
        for cell in row:
            if cell.value == "· 1행 = 폼 1건 제출. 실행=Y 행만 처리.":
                cell.value = "· 1행 = 폼 1건. 실행=Y 행을 검토하고, 명시적인 제출 승인 후 전송."

    # Reference examples must not retain the original customer's identity.
    substitutions = sorted(
        (
            (value, replacements.get(label, "[example]"))
            for label, value in source_values.items()
            if value
        ),
        key=lambda pair: len(pair[0]),
        reverse=True,
    )
    replacement_count = 0
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                cell.comment = None
                if not isinstance(cell.value, str) or cell.data_type == "f":
                    continue
                text = cell.value
                for old, new in substitutions:
                    text = text.replace(old, new)
                if text != cell.value:
                    cell.value = text
                    replacement_count += 1
    workbook.properties.creator = "Foundry Quota Increase Skill"
    workbook.properties.lastModifiedBy = "Foundry Quota Increase Skill"
    workbook.properties.title = "Microsoft Foundry quota increase request template"
    workbook.properties.subject = "Blank request inputs and reference-only form choices"
    workbook.properties.description = (
        "Request inputs and common values are blank. "
        "Execution=Y selects rows but does not replace human submission approval."
    )
    for attribute in ("keywords", "category", "identifier", "contentStatus", "revision"):
        setattr(workbook.properties, attribute, None)
    workbook.custom_doc_props = CustomPropertyList()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = False
    workbook.calculation.calcMode = "auto"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    sanitize_package_metadata(output)

    check = load_workbook(output)
    request_check = find_sheet(check, "요청목록")
    common_check = find_sheet(check, "공통정보")
    assert all(
        request_check.cell(row, column).value is None
        for row in range(4, request_check.max_row + 1)
        for column in (1, 2, 3, 4, 5, 6, 7, 8, 10, 11)
    )
    assert all(
        common_check.cell(row, 2).value is None
        for row in range(2, common_check.max_row + 1)
    )
    assert len(request_check.data_validations.dataValidation) == validation_count
    assert len(check.defined_names) == defined_name_count
    check.close()
    return {
        "status": "success", "output": str(output),
        "request_input_rows_cleared": requests.max_row - 3,
        "common_values_cleared": common.max_row - 1,
        "reference_examples_sanitized": replacement_count,
        "reference_sheets": [reference.title, choices.title],
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    print(json.dumps(make_template(args.source, args.output), ensure_ascii=False))

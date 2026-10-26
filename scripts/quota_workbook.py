#!/usr/bin/env python3
"""Read-only quota intake and a human-approved, fail-closed submission ledger."""

import argparse
from contextlib import contextmanager
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import time
import uuid
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


SCHEMA_VERSION = 1
BROWSER_TEMPLATE = Path(__file__).with_name("browser_flow.js")
COMMON_FIELDS = (
    "First Name", "Last Name", "Company Email", "Company Name",
    "Company Address", "Company City", "Company Postal Code",
    "Company Country", "Subscription Id", "Justification",
)
REQUEST_FIELDS = (
    "request_id", "execute", "quota_type", "model_type", "deployment_type",
    "model_name", "region", "requested_tpm", "quota_input", "subscription_id",
)
INPUT_FIELDS = (
    "execute", "request_id", "quota_type", "model_type", "deployment_type",
    "model_name", "region", "requested_tpm", "quota_input", "subscription_id",
    "notes",
)
WORKBOOK_HEADERS = {
    "실행(Y/N)": "execute",
    "요청ID": "request_id",
    "Quota Type(폼 11번)": "quota_type",
    "Model Type(폼 12번)": "model_type",
    "배포 유형(폼 13번 Quota Request Type)": "deployment_type",
    "폼 모델명(정확한 표기)": "model_name",
    "Region / Data Zone Pool": "region",
    "신청 총 TPM(기존+추가)": "requested_tpm",
    "폼 입력값 kTPM(자동)": "quota_input",
    "구독 ID(행별 지정 시)": "subscription_id",
    "비고": "notes",
}
OPENAI_TYPES = {"Azure OpenAI", "Azure OpenAI Service"}
ADM_TYPES = {
    "Azure Direct Model (ADM)", "Azure Direct Models (ADM)",
    "Azure Direct Model", "Azure Direct Models",
}
OPENAI_DEPLOYMENTS = {"Global Standard"}
MAX_SAFE_INTEGER = 9007199254740991
ZERO_GUID = "00000000-0000-0000-0000-000000000000"
GUID = re.compile(r"^[0-9a-fA-F]{8}(?:-[0-9a-fA-F]{4}){3}-[0-9a-fA-F]{12}$")
EMAIL = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
UNSUPPORTED_OPENAI = re.compile(
    r"audio|realtime|real[- _]?time|live|whisper|tts|transcri(?:be|pt)|translat|"
    r"embedding|dall[- _]?e|image|video|sora",
    re.IGNORECASE,
)
TERMINAL_OR_ACTIVE = {"submitting", "submitted", "unknown"}


class LedgerError(ValueError):
    pass


def utc_now():
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def canonical(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=False,
                      separators=(",", ":"), allow_nan=False)


def digest(value):
    return hashlib.sha256(canonical(value).encode("utf-8")).hexdigest()


def read_json(path):
    with Path(path).open("r", encoding="utf-8-sig") as stream:
        value = json.load(stream)
    if not isinstance(value, dict):
        raise LedgerError("JSON root must be an object.")
    return value


def atomic_json(path, value, exclusive=False):
    path = Path(path)
    staged = path.with_name("." + path.name + "." + uuid.uuid4().hex + ".tmp")
    try:
        fd = os.open(staged, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False))
            stream.write("\n")
            stream.flush()
            os.fsync(stream.fileno())
        if exclusive:
            # Linking the complete file makes publication both atomic and no-clobber.
            os.link(staged, path)
        else:
            for attempt in range(6):
                try:
                    os.replace(staged, path)
                    break
                except PermissionError:
                    if attempt == 5:
                        raise
                    # Windows scanners/sync clients can briefly hold the replaced file.
                    time.sleep(0.05 * (attempt + 1))
    finally:
        if staged.exists():
            staged.unlink()


def manifest_path(run_dir):
    return Path(run_dir).resolve() / "manifest.json"


def load_manifest(run_dir):
    manifest = read_json(manifest_path(run_dir))
    if manifest.get("schema_version") != SCHEMA_VERSION:
        raise LedgerError("Unsupported manifest schema.")
    if manifest.get("mode") not in {"capture", "submit"}:
        raise LedgerError("Invalid manifest mode.")
    if not isinstance(manifest.get("common"), dict) or not isinstance(manifest.get("requests"), list):
        raise LedgerError("Malformed manifest.")
    return manifest


@contextmanager
def mutation(run_dir):
    directory = Path(run_dir).resolve()
    lock = directory / ".ledger.lock"
    try:
        fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError as exc:
        raise LedgerError("Run is locked; reconcile the active process before retrying.") from exc
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            stream.write(str(os.getpid()))
        manifest = load_manifest(directory)
        yield manifest
        atomic_json(manifest_path(directory), manifest)
    finally:
        lock.unlink()


def event(manifest, kind, request_id=None, **fields):
    item = {"type": kind, "timestamp": utc_now(), **fields}
    if request_id is not None:
        item["request_id"] = request_id
    manifest["events"].append(item)


def text(value):
    if value is None:
        return ""
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise LedgerError("Input values must be text or scalar numbers, not booleans.")
    return str(value).strip()


def input_text(value, label, issues):
    try:
        result = text(value)
    except LedgerError:
        issues.append(label + ": expected text, not a boolean or structured value.")
        return ""
    if result.startswith("="):
        issues.append(label + ": formulas are not permitted in input cells.")
        return ""
    return result


def normalize_common(raw, extra_issues=None):
    if not isinstance(raw, dict):
        raise LedgerError("common must be an object with original form field labels.")
    issues = list(extra_issues or [])
    common = {}
    for key, value in raw.items():
        if not isinstance(key, str):
            raise LedgerError("Common field labels must be strings.")
        label = key.strip()
        if label in common:
            issues.append("common: duplicate field label.")
        common[label] = input_text(value, "common " + label, issues)
    for label in COMMON_FIELDS:
        common.setdefault(label, "")
        if label != "Subscription Id" and not common[label]:
            issues.append("common " + label + ": required.")
    if common["Company Email"] and not EMAIL.fullmatch(common["Company Email"]):
        issues.append("common Company Email: invalid email structure.")
    return common, issues


def parse_tpm(value, issues):
    if isinstance(value, str):
        if value.strip().startswith("="):
            issues.append("requested_tpm: formulas are not permitted.")
            return None
        if not re.fullmatch(r"[0-9]+", value.strip()):
            issues.append("requested_tpm: use an integer without commas or fractions.")
            return None
        value = int(value.strip())
    if isinstance(value, bool) or not isinstance(value, int):
        issues.append("requested_tpm: must be an integer, not a float or boolean.")
        return None
    if value > MAX_SAFE_INTEGER:
        issues.append("requested_tpm: exceeds JavaScript Number.MAX_SAFE_INTEGER; cannot render exactly.")
        return None
    if value <= 0 or value % 1000:
        issues.append("requested_tpm: must be positive and an exact multiple of 1000.")
        return None
    return value


def validate_branch(row, issues):
    if row["quota_type"] != "Model Deployment (PTU/RPM/TPM)":
        issues.append("quota_type: only Model Deployment (PTU/RPM/TPM) is supported.")
    model_type = row["model_type"]
    if model_type in OPENAI_TYPES:
        if row["deployment_type"] not in OPENAI_DEPLOYMENTS:
            issues.append("deployment_type: Azure OpenAI supports only Global Standard chat TPM "
                          "in v1; Standard, Data Zone Standard, RPM, and PTU are unsupported.")
        if UNSUPPORTED_OPENAI.search(row["model_name"]):
            issues.append("model_name: non-chat/audio/realtime/live/speech models are unsupported by v1 TPM intake.")
    elif model_type in ADM_TYPES:
        if row["deployment_type"] != "Global Standard":
            issues.append("deployment_type: ADM supports only Global Standard in v1.")
    else:
        issues.append("model_type: only Azure OpenAI chat TPM or Azure Direct Model (ADM) is supported.")
    for field in ("quota_type", "model_type", "deployment_type", "model_name"):
        if not row[field]:
            issues.append(field + ": required; use the exact live form spelling.")


def normalize_request(raw, source_row, common, common_issues, generated_id=None):
    if not isinstance(raw, dict):
        raise LedgerError("Every request must be an object.")
    issues = list(common_issues)
    row = {field: input_text(raw.get(field), field, issues)
           for field in INPUT_FIELDS if field not in {"requested_tpm", "quota_input"}}
    if "request_id" not in raw and generated_id is not None:
        row["request_id"] = generated_id
    row["execute"] = row["execute"].upper() or "N"
    if row["execute"] not in {"Y", "N"}:
        issues.append("execute: must be Y or N.")
    if not row["request_id"]:
        issues.append("request_id: required and must be unique.")
    row["source_row"] = source_row
    row["region"] = row["region"] or None
    row["requested_tpm"] = parse_tpm(raw.get("requested_tpm"), issues)
    row["quota_input"] = (str(row["requested_tpm"] // 1000)
                          if row["requested_tpm"] is not None else "")
    row["subscription_id"] = row["subscription_id"] or common["Subscription Id"]
    if not GUID.fullmatch(row["subscription_id"]) or row["subscription_id"].lower() == ZERO_GUID:
        issues.append("subscription_id: a nonzero GUID is required; all-zero placeholders are "
                      "invalid (row value overrides common).")
    validate_branch(row, issues)
    row.update(issues=issues, ui_issues=[], status="pending", selected=False,
               selection_reason="", ui_evidence=None, attempts=[])
    row["payload_hash"] = payload_hash(common, row)
    return row


def request_payload(row):
    return {field: row.get(field) for field in REQUEST_FIELDS}


def payload_hash(common, row):
    return digest({"common": common, "request": request_payload(row)})


def current_payload_valid(manifest, row):
    if row.get("payload_hash") != payload_hash(manifest["common"], row):
        raise LedgerError("Immutable payload changed; create a new reviewed run.")
    normalized_common, common_issues = normalize_common(manifest["common"])
    fresh = normalize_request(row, row.get("source_row"), normalized_common, common_issues)
    if fresh["issues"] or request_payload(fresh) != request_payload(row):
        raise LedgerError("Current payload is invalid; create a new reviewed run.")


def select_rows(rows, mode, include_disabled, ids):
    requested_ids = list(ids or [])
    if len(set(requested_ids)) != len(requested_ids):
        raise LedgerError("--ids must not contain duplicate IDs.")
    if set(requested_ids) - {row["request_id"] for row in rows}:
        raise LedgerError("One or more requested IDs do not exist in the input.")
    for row in rows:
        wanted = not requested_ids or row["request_id"] in requested_ids
        enabled = row["execute"] == "Y"
        selected = wanted and (enabled or (mode == "capture" and
                                            (include_disabled or bool(requested_ids))))
        row["selected"] = selected
        row["selection_reason"] = (
            "selected enabled row" if selected and enabled else
            "explicit capture-only disabled row" if selected else
            "outside --ids scope" if not wanted else
            "disabled rows are never eligible for submission" if mode == "submit" else
            "disabled row; capture requires explicit inclusion"
        )
    by_id = {}
    for row in rows:
        if row["request_id"]:
            by_id.setdefault(row["request_id"], []).append(row)
    for duplicates in by_id.values():
        if len(duplicates) > 1:
            for row in duplicates:
                row["issues"].append("request_id: duplicated in input; assign unique IDs.")
    duplicates = {}
    for row in rows:
        if row["selected"] and not row["issues"]:
            body = request_payload(row)
            body.pop("request_id")
            body.pop("execute")
            duplicates.setdefault(digest(body), []).append(row)
    for group in duplicates.values():
        if len(group) > 1:
            for row in group:
                row["issues"].append("payload: exact duplicate of another selected request.")
    for row in rows:
        row["status"] = ("invalid" if row["issues"] else "pending") if row["selected"] else "skipped"


def finish_prepare(source, common_raw, raw_rows, run_dir, mode="capture",
                   include_disabled=False, ids=None, common_issues=None, interactive=False):
    if mode not in {"capture", "submit"}:
        raise LedgerError("Mode must be capture or submit.")
    common, issues = normalize_common(common_raw, common_issues)
    rows = [normalize_request(raw, source_row, common, issues,
                              f"REQ-{index:03d}" if interactive else None)
            for index, (source_row, raw) in enumerate(raw_rows, 1)]
    select_rows(rows, mode, include_disabled, ids)
    manifest = {
        "schema_version": SCHEMA_VERSION, "run_id": str(uuid.uuid4()), "mode": mode, "source": source,
        "created_at": utc_now(), "common": common, "requests": rows,
        "selection": {"include_disabled": bool(include_disabled), "ids": list(ids or []),
                      "execute_y_is_consent": False},
        "approval": None, "preview": None, "events": [],
        "input_schema": {"common": list(COMMON_FIELDS), "requests": list(INPUT_FIELDS),
                         "quota_unit": "H/requested_tpm is total TPM; quota_input = TPM / 1000"},
    }
    event(manifest, "prepared")
    directory = Path(run_dir).resolve()
    if directory.exists():
        raise LedgerError("Run directory already exists; use a NEW run directory.")
    directory.mkdir(parents=True, exist_ok=False, mode=0o700)
    atomic_json(directory / "manifest.json", manifest, exclusive=True)
    return manifest


def sheet(workbook, names):
    normalized = {"".join(name.split()) for name in names}
    matches = [ws for ws in workbook.worksheets if "".join(ws.title.split()) in normalized]
    if len(matches) != 1:
        raise LedgerError("Required intake sheet missing or ambiguous: " + names[0])
    return matches[0]


def prepare(workbook, run_dir, mode="capture", include_disabled=False, ids=None):
    path = Path(workbook).resolve()
    before_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    common = {}
    common_issues = []
    raw_rows = []
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        common_sheet = sheet(book, ("공통정보",))
        for cells in common_sheet.iter_rows(min_col=1, max_col=2):
            label = cells[0].value
            if label is None or str(label).strip() in {"항목", ""}:
                continue
            label = str(label).strip()
            if cells[0].data_type == "f":
                common_issues.append("common: formulas are not permitted in field labels.")
                continue
            if label in common:
                common_issues.append("common: duplicate field label.")
            common[label] = cells[1].value
        requests_sheet = sheet(book, ("요청목록",))
        header_cells = next(requests_sheet.iter_rows(min_row=3, max_row=3), ())
        columns = {}
        for index, cell in enumerate(header_cells):
            field = WORKBOOK_HEADERS.get(cell.value) if isinstance(cell.value, str) else None
            if field is not None:
                if field in columns:
                    raise LedgerError("Duplicate request header in row 3; each input header must be unique.")
                columns[field] = index
        for header, field in WORKBOOK_HEADERS.items():
            if field not in columns:
                raise LedgerError("Missing exact request header in row 3: " + header)
        for row_number, cells in enumerate(
                requests_sheet.iter_rows(min_row=4, min_col=1, max_col=max(columns.values()) + 1), 4):
            raw = {field: cells[index].value for field, index in columns.items()}
            meaningful = [value for field, value in raw.items()
                          if field not in {"execute", "quota_input"}]
            if all(value is None or value == "" for value in meaningful) and (
                    raw["execute"] is None or str(raw["execute"]).strip().upper() in {"", "N"}):
                continue
            raw_rows.append((row_number, raw))
    finally:
        book.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != before_hash:
        raise LedgerError("Source workbook changed during intake; retry with a stable copy.")
    return finish_prepare({"kind": "xlsx", "path": str(path), "sha256": before_hash},
                          common, raw_rows, run_dir, mode, include_disabled, ids, common_issues)


def prepare_json(input_file, run_dir, mode="capture", include_disabled=False, ids=None):
    path = Path(input_file).resolve()
    source_bytes = path.read_bytes()
    data = read_json(path)
    if not isinstance(data.get("requests"), list):
        raise LedgerError("requests must be an array.")
    return finish_prepare(
        {"kind": "json", "path": str(path), "sha256": hashlib.sha256(source_bytes).hexdigest()},
        data.get("common"), [(None, row) for row in data["requests"]],
        run_dir, mode, include_disabled, ids, interactive=True,
    )


def find_row(manifest, request_id):
    rows = [row for row in manifest["requests"] if row.get("request_id") == request_id]
    if len(rows) != 1:
        raise LedgerError("Request ID is missing or ambiguous.")
    return rows[0]


def png_exists(value):
    if not isinstance(value, str):
        return False
    path = Path(value)
    return path.is_absolute() and path.suffix.lower() == ".png" and path.is_file()


def ui_problems(manifest, row, evidence):
    issues = []
    if evidence.get("request_id") != row["request_id"] or evidence.get("payload_hash") != row["payload_hash"]:
        issues.append("UI evidence ID/hash does not match planned payload.")
    if evidence.get("status") != "filled":
        issues.append("Live form is blocked: " + str(evidence.get("status", "missing status")))
    for field in ("quota_type", "model_type", "deployment_type", "model_name", "quota_input"):
        if evidence.get(field) != row[field]:
            issues.append("UI field does not exactly match plan: " + field)
    if evidence.get("common_verified") is not True:
        issues.append("Common fields were not all verified.")
    if evidence.get("subscription_id_verified") is not True:
        issues.append("Resolved subscription was not verified.")
    if "subscription_id" in evidence and evidence["subscription_id"] != row["subscription_id"]:
        issues.append("UI subscription does not match plan.")
    if not isinstance(evidence.get("region_asked"), bool):
        issues.append("region_asked must explicitly be true or false.")
    elif evidence["region_asked"] and (
            not row["region"] or evidence.get("region_selected") != row["region"]):
        issues.append("Live form requires a matching nonblank planned region; prepare a new run if missing.")
    elif not evidence["region_asked"] and evidence.get("region_selected") is not None:
        issues.append("Region was not asked but evidence reports a selection.")
    screenshots = evidence.get("screenshot_paths")
    if not isinstance(screenshots, list) or not screenshots or not all(png_exists(p) for p in screenshots):
        issues.append("At least one absolute existing local PNG screenshot is required.")
    if any(evidence.get(field) for field in ("visible_errors", "errors", "validation_errors")):
        issues.append("Live form reports visible or validation errors.")
    return issues


def record_ui(run_dir, request_id, evidence_file):
    evidence = read_json(evidence_file)
    with mutation(run_dir) as manifest:
        row = find_row(manifest, request_id)
        if row["status"] in TERMINAL_OR_ACTIVE:
            raise LedgerError("Submission state cannot be overwritten by UI evidence.")
        if not row["selected"] or row["issues"]:
            raise LedgerError("Only selected, valid intake rows can record UI evidence.")
        current_payload_valid(manifest, row)
        row["ui_issues"] = ui_problems(manifest, row, evidence)
        row["ui_evidence"] = evidence
        row["status"] = "blocked" if row["ui_issues"] else "filled"
        event(manifest, "ui_recorded", request_id, status=row["status"])
        return {"request_id": request_id, "status": row["status"], "issues": row["ui_issues"]}


def is_ready(manifest, row):
    if not row.get("selected") or row["execute"] != "Y" or row["status"] != "filled" or row["issues"]:
        return False
    try:
        current_payload_valid(manifest, row)
    except LedgerError:
        return False
    return isinstance(row.get("ui_evidence"), dict) and not ui_problems(manifest, row, row["ui_evidence"])


def region_facts(row):
    evidence = row.get("ui_evidence") or {}
    asked = evidence.get("region_asked")
    selected = evidence.get("region_selected")
    effective = "not asked" if asked is False else selected if asked is True else "unverified"
    return {"region_asked": asked, "region_selected": selected, "effective_region": effective}


def scope_fingerprint(manifest, request_ids):
    return digest({"common": manifest["common"],
                   "requests": [{**request_payload(find_row(manifest, request_id)),
                                 **region_facts(find_row(manifest, request_id))}
                                for request_id in request_ids]})


def preview(run_dir):
    with mutation(run_dir) as manifest:
        ready_ids = [row["request_id"] for row in manifest["requests"] if is_ready(manifest, row)]
        fingerprint = scope_fingerprint(manifest, ready_ids)
        rows = []
        for row in manifest["requests"]:
            classification = ("skipped" if not row["selected"] or row["execute"] != "Y" else
                              "ready" if row["request_id"] in ready_ids else "blocked")
            issues = list(row["issues"]) + list(row.get("ui_issues", []))
            if classification == "blocked" and not issues:
                issues.append("Not ready: current state is " + row["status"] + ".")
            rows.append({**request_payload(row), **region_facts(row), "source_row": row["source_row"],
                         "payload_hash": row["payload_hash"], "status": row["status"],
                         "classification": classification, "issues": issues,
                         "selection_reason": row["selection_reason"]})
        result = {"schema_version": 1, "mode": manifest["mode"], "common": manifest["common"],
                  "requests": rows, "ready_request_ids": ready_ids,
                  "preview_fingerprint": fingerprint,
                  "notice": "Y is eligibility, not consent. Approval covers only the displayed ready IDs; "
                            "blocked and skipped rows are outside approval scope."}
        manifest["preview"] = {"fingerprint": fingerprint, "request_ids": ready_ids,
                               "timestamp": utc_now()}
        event(manifest, "previewed", request_ids=ready_ids, fingerprint=fingerprint)
        return result


def approve(run_dir, fingerprint, confirmation_text):
    if not isinstance(confirmation_text, str) or not confirmation_text.strip():
        raise LedgerError("Nonblank explicit human confirmation text is required.")
    with mutation(run_dir) as manifest:
        if manifest["mode"] != "submit":
            raise LedgerError("Capture runs cannot be approved for submission.")
        if manifest["approval"] is not None:
            raise LedgerError("This run already has an approval; scope cannot expand. Use a new run.")
        shown = manifest.get("preview")
        ready_ids = [row["request_id"] for row in manifest["requests"] if is_ready(manifest, row)]
        if not shown or not ready_ids or shown["request_ids"] != ready_ids:
            raise LedgerError("Preview the current nonempty ready scope before approval.")
        if fingerprint != shown["fingerprint"] or fingerprint != scope_fingerprint(manifest, ready_ids):
            raise LedgerError("Preview fingerprint is stale or incorrect.")
        manifest["approval"] = {
            "request_ids": ready_ids, "fingerprint": fingerprint,
            "run_id": manifest["run_id"], "timestamp": utc_now(),
            "confirmation_text": confirmation_text.strip(),
        }
        event(manifest, "approved", request_ids=ready_ids, fingerprint=fingerprint)
        return {"approved_request_ids": ready_ids, "fingerprint": fingerprint}


def validate_approval(manifest, row):
    approval = manifest.get("approval")
    if manifest["mode"] != "submit" or row["execute"] != "Y" or not row["selected"]:
        raise LedgerError("Only selected Y rows in submit mode may submit.")
    if not approval or row["request_id"] not in approval.get("request_ids", []):
        raise LedgerError("Request has no explicit human approval for this exact scope.")
    if approval.get("run_id") != manifest["run_id"]:
        raise LedgerError("Approval belongs to another run.")
    approved_ids = approval["request_ids"]
    if not approved_ids or len(set(approved_ids)) != len(approved_ids):
        raise LedgerError("Invalid approval scope.")
    for request_id in approved_ids:
        approved_row = find_row(manifest, request_id)
        if approved_row["execute"] != "Y" or not approved_row["selected"]:
            raise LedgerError("Approved scope changed.")
        current_payload_valid(manifest, approved_row)
    if not approval.get("confirmation_text", "").strip() or (
            scope_fingerprint(manifest, approved_ids) != approval["fingerprint"]):
        raise LedgerError("Approval is stale; immutable scope or common values changed.")
    return approval


def active_attempt(row):
    attempts = row.get("attempts", [])
    if not attempts:
        raise LedgerError("No persisted submission attempt.")
    return attempts[-1]


def begin_submit(run_dir, request_id, output=None):
    ticket = None
    with mutation(run_dir) as manifest:
        row = find_row(manifest, request_id)
        approval = validate_approval(manifest, row)
        if any(other["status"] in {"submitting", "unknown"} for other in manifest["requests"]):
            raise LedgerError("Another submitting/unknown attempt requires reconciliation; stop.")
        if not is_ready(manifest, row):
            raise LedgerError("Request must be filled and verified, not already submitted.")
        attempt = {"attempt_id": str(uuid.uuid4()), "timestamp": utc_now(), "status": "submitting"}
        ticket_path = (Path(output).resolve() if output else
                       Path(run_dir).resolve() / ("ticket-" + attempt["attempt_id"] + ".json"))
        if ticket_path.exists() or not ticket_path.parent.is_dir():
            raise LedgerError("Ticket destination must be new and its parent directory must exist.")
        ticket = {
            "schema_version": 1, "approved": True, "request_id": request_id,
            "payload_hash": row["payload_hash"], "approval_fingerprint": approval["fingerprint"],
            "attempt_id": attempt["attempt_id"], "timestamp": attempt["timestamp"],
            "run_dir": str(Path(run_dir).resolve()), "common": manifest["common"],
            "request": request_payload(row),
        }
        attempt["ticket_hash"] = digest(ticket)
        attempt["ticket_path"] = str(ticket_path)
        row["attempts"].append(attempt)
        row["status"] = "submitting"
        event(manifest, "submission_begun", request_id, attempt_id=attempt["attempt_id"])
    # The durable stop guard must precede publication of a click-capable ticket.
    atomic_json(ticket_path, ticket, exclusive=True)
    return {"request_id": request_id, "status": "submitting", "ticket_path": str(ticket_path),
            "attempt_id": ticket["attempt_id"]}


def recognized_success(value):
    if not isinstance(value, str):
        return False
    return bool(re.fullmatch(
        r"(?:Your response (?:was|has been) submitted|응답이 제출되었습니다|응답을 제출했습니다|Thanks! You've completed the request!)[.!]?",
        value.strip(), flags=re.IGNORECASE,
    ))


def finish_submit(run_dir, request_id, evidence_file):
    evidence = read_json(evidence_file)
    with mutation(run_dir) as manifest:
        row = find_row(manifest, request_id)
        if row["status"] != "submitting":
            raise LedgerError("Only a currently submitting attempt can be finished.")
        attempt = active_attempt(row)
        if (evidence.get("request_id") != request_id or
                evidence.get("attempt_id") != attempt["attempt_id"] or
                evidence.get("payload_hash") != row["payload_hash"]):
            raise LedgerError("Submission evidence does not match the persisted attempt.")
        current_payload_valid(manifest, row)
        claimed = evidence.get("submission_status")
        clicked = evidence.get("submit_clicked")
        status = "unknown"
        if (claimed == "submitted" and clicked is True and
                recognized_success(evidence.get("confirmation_text")) and
                png_exists(evidence.get("confirmation_screenshot")) and
                not any(evidence.get(k) for k in ("visible_errors", "errors", "validation_errors"))):
            status = "submitted"
        elif claimed == "not_clicked" and clicked is False and (
                isinstance(evidence.get("detail"), str) and evidence["detail"].strip()):
            status = "filled"
        row["status"] = status
        attempt["status"] = "not_clicked" if status == "filled" else status
        attempt["finished_at"] = utc_now()
        attempt["evidence"] = evidence
        event(manifest, "submission_finished", request_id, status=status,
              attempt_id=attempt["attempt_id"])
        return {"request_id": request_id, "status": status,
                "stop_and_reconcile": status == "unknown"}


def mark_unknown(run_dir, request_id, detail):
    if not isinstance(detail, str) or not detail.strip():
        raise LedgerError("Explicit timeout/failure detail is required.")
    with mutation(run_dir) as manifest:
        row = find_row(manifest, request_id)
        if row["status"] != "submitting":
            raise LedgerError("Only a submitting attempt may be marked unknown.")
        attempt = active_attempt(row)
        attempt.update(status="unknown", detail=detail.strip(), finished_at=utc_now())
        row["status"] = "unknown"
        event(manifest, "submission_unknown", request_id, attempt_id=attempt["attempt_id"])
        return {"request_id": request_id, "status": "unknown", "stop_and_reconcile": True}


def reconcile_submitted(run_dir, request_id, evidence_file, confirmation_text):
    if not isinstance(confirmation_text, str) or not confirmation_text.strip():
        raise LedgerError("Explicit human confirmation of the actual receipt is required.")
    evidence = read_json(evidence_file)
    with mutation(run_dir) as manifest:
        row = find_row(manifest, request_id)
        if row["status"] != "unknown":
            raise LedgerError("Only an unknown attempt may be reconciled; never resubmit it.")
        validate_approval(manifest, row)
        attempt = active_attempt(row)
        if attempt["status"] != "unknown" or not attempt.get("rendered_at"):
            raise LedgerError("Reconciliation requires an existing rendered unknown attempt.")
        if (evidence.get("request_id") != request_id or
                evidence.get("attempt_id") != attempt["attempt_id"] or
                evidence.get("payload_hash") != row["payload_hash"]):
            raise LedgerError("Reconciliation evidence does not match the persisted attempt.")
        if (evidence.get("submit_clicked") is not True or
                evidence.get("submission_status") != "submitted" or
                not recognized_success(evidence.get("confirmation_text")) or
                not png_exists(evidence.get("confirmation_screenshot")) or
                any(evidence.get(k) for k in ("visible_errors", "errors", "validation_errors"))):
            raise LedgerError("Explicit completion text and screenshot without errors are required.")
        previous = attempt.get("evidence", {})
        if previous.get("submit_clicked") is False:
            raise LedgerError("Existing evidence says Submit was not clicked; reconcile manually.")
        attempt["reconciliation"] = {
            "timestamp": utc_now(), "confirmation_text": confirmation_text.strip(),
            "evidence": evidence,
        }
        attempt["status"] = row["status"] = "submitted"
        attempt["finished_at"] = utc_now()
        event(manifest, "submission_reconciled", request_id, attempt_id=attempt["attempt_id"])
        return {"request_id": request_id, "status": "submitted", "new_clicks": 0,
                "stop_and_reconcile": any(other["status"] in {"submitting", "unknown"}
                                          for other in manifest["requests"])}


def validate_ticket(manifest, row, ticket, run_dir):
    approval = validate_approval(manifest, row)
    attempt = active_attempt(row)
    if row["status"] != "submitting" or attempt["status"] != "submitting":
        raise LedgerError("Submit render requires the persisted active submitting attempt.")
    if any(other["status"] == "unknown" or (
            other is not row and other["status"] == "submitting") for other in manifest["requests"]):
        raise LedgerError("Unreconciled attempt prevents submission.")
    if (ticket.get("approved") is not True or ticket.get("request_id") != row["request_id"] or
            ticket.get("attempt_id") != attempt["attempt_id"] or
            ticket.get("payload_hash") != row["payload_hash"] or
            ticket.get("approval_fingerprint") != approval["fingerprint"] or
            ticket.get("run_dir") != str(Path(run_dir).resolve()) or
            ticket.get("common") != manifest["common"] or
            ticket.get("request") != request_payload(row) or digest(ticket) != attempt["ticket_hash"]):
        raise LedgerError("Submit ticket does not match the exact persisted attempt and scope.")
    if attempt.get("rendered_at"):
        raise LedgerError("This attempt already produced a submit script; never replay it. Reconcile first.")
    return attempt


def render(run_dir, request_id, action, output, screenshots_dir, ticket_file=None):
    if action not in {"fill", "submit"}:
        raise LedgerError("Render action must be fill or submit.")
    path = Path(output).resolve()
    screenshot_directory = Path(screenshots_dir).resolve()
    if path.exists() or not path.parent.is_dir():
        raise LedgerError("Render output must be new and its parent directory must exist.")
    if path == screenshot_directory:
        raise LedgerError("Render output and screenshot directory must be different paths.")
    template = BROWSER_TEMPLATE.read_text(encoding="utf-8")
    if template.count("__QUOTA_CONFIG__") != 1:
        raise LedgerError("Browser template must contain __QUOTA_CONFIG__ exactly once.")
    ticket = read_json(ticket_file) if ticket_file else None
    with mutation(run_dir) as manifest:
        row = find_row(manifest, request_id)
        current_payload_valid(manifest, row)
        attempt = None
        if action == "fill":
            if ticket_file or not row["selected"] or row["issues"] or row["status"] in TERMINAL_OR_ACTIVE:
                raise LedgerError("Fill requires a selected valid row without an active/terminal attempt.")
        else:
            if ticket is None:
                raise LedgerError("Submit rendering requires a persisted approval ticket.")
            attempt = validate_ticket(manifest, row, ticket, run_dir)
        if screenshot_directory.exists() and not screenshot_directory.is_dir():
            raise LedgerError("Screenshot destination exists but is not a directory.")
        screenshot_directory.mkdir(parents=True, exist_ok=True, mode=0o700)
        if attempt is not None:
            attempt["rendered_at"] = utc_now()
        config = {"action": action, "common": manifest["common"],
                  "request": {**request_payload(row), "source_row": row["source_row"],
                              "payload_hash": row["payload_hash"], "status": row["status"],
                              "issues": list(row["issues"])},
                  "mode": manifest["mode"], "ticket": ticket,
                  "screenshot_dir": str(screenshot_directory)}
        encoded = json.dumps(canonical(config), ensure_ascii=True)
        script = template.replace("__QUOTA_CONFIG__", "JSON.parse(" + encoded + ")")
        event(manifest, "script_rendered", request_id, action=action)
    staged = path.with_name("." + path.name + "." + uuid.uuid4().hex + ".tmp")
    try:
        fd = os.open(staged, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(script)
            stream.flush()
            os.fsync(stream.fileno())
        os.link(staged, path)
    finally:
        if staged.exists():
            staged.unlink()
    return {"request_id": request_id, "action": action, "output": str(path),
            "notice": "Rendering did not execute a browser. Keep private; never replay a submit script."}


def status(run_dir):
    manifest = load_manifest(run_dir)
    counts = {}
    for row in manifest["requests"]:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    return {"schema_version": 1, "mode": manifest["mode"], "counts": counts,
            "total_requests": len(manifest["requests"]),
            "selected_count": sum(bool(row["selected"]) for row in manifest["requests"]),
            "submitted_count": counts.get("submitted", 0),
            "approved_request_ids": (manifest["approval"] or {}).get("request_ids", []),
            "stop_and_reconcile": any(row["status"] in {"submitting", "unknown"}
                                      for row in manifest["requests"])}


def parser():
    root = argparse.ArgumentParser(description=__doc__)
    commands = root.add_subparsers(dest="command", required=True)
    for name, source_flag in (("prepare", "--workbook"), ("prepare-json", "--input")):
        command = commands.add_parser(name)
        command.add_argument(source_flag, required=True)
        command.add_argument("--run-dir", required=True)
        command.add_argument("--mode", choices=("capture", "submit"), default="capture")
        command.add_argument("--include-disabled", action="store_true")
        command.add_argument("--ids", nargs="+")
    for name in ("record-ui", "preview", "approve", "begin-submit", "finish-submit",
                 "mark-unknown", "reconcile-submitted", "status", "render"):
        command = commands.add_parser(name)
        command.add_argument("--run-dir", required=True)
        if name in {"record-ui", "begin-submit", "finish-submit", "mark-unknown",
                    "reconcile-submitted", "render"}:
            command.add_argument("--id", required=True)
        if name in {"record-ui", "finish-submit", "reconcile-submitted"}:
            command.add_argument("--evidence", required=True)
        if name == "approve":
            command.add_argument("--fingerprint", required=True)
        if name in {"approve", "reconcile-submitted"}:
            command.add_argument("--confirmation-text", required=True)
        if name == "begin-submit":
            command.add_argument("--output")
        if name == "mark-unknown":
            command.add_argument("--detail", required=True)
        if name == "render":
            command.add_argument("--action", choices=("fill", "submit"), required=True)
            command.add_argument("--output", required=True)
            command.add_argument("--screenshots-dir", required=True)
            command.add_argument("--ticket")
    return root


def main(argv=None):
    args = parser().parse_args(argv)
    try:
        if args.command in {"prepare", "prepare-json"}:
            function = prepare if args.command == "prepare" else prepare_json
            source = args.workbook if args.command == "prepare" else args.input
            function(source, args.run_dir, args.mode, args.include_disabled, args.ids)
            result = {**status(args.run_dir), "manifest": str(manifest_path(args.run_dir))}
        elif args.command == "record-ui":
            result = record_ui(args.run_dir, args.id, args.evidence)
        elif args.command == "preview":
            result = preview(args.run_dir)
        elif args.command == "approve":
            result = approve(args.run_dir, args.fingerprint, args.confirmation_text)
        elif args.command == "begin-submit":
            result = begin_submit(args.run_dir, args.id, args.output)
        elif args.command == "finish-submit":
            result = finish_submit(args.run_dir, args.id, args.evidence)
        elif args.command == "mark-unknown":
            result = mark_unknown(args.run_dir, args.id, args.detail)
        elif args.command == "reconcile-submitted":
            result = reconcile_submitted(args.run_dir, args.id, args.evidence,
                                         args.confirmation_text)
        elif args.command == "render":
            result = render(args.run_dir, args.id, args.action, args.output,
                            args.screenshots_dir, args.ticket)
        else:
            result = status(args.run_dir)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except LedgerError as exc:
        print("Error: " + str(exc), file=sys.stderr)
        return 2
    except (ValueError, OSError, BadZipFile, InvalidFileException, KeyError, TypeError):
        # Do not echo parser/library exceptions that can contain private input values.
        print("Error: intake or ledger operation rejected. Check input/schema, "
              "row issues, approval scope, evidence, paths, and active-attempt guards.",
              file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())

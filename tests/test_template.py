from pathlib import Path
from xml.dom import minidom
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
import pytest

import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from make_blank_template import make_template, sanitize_package_metadata


def synthetic_source(path):
    workbook = Workbook()
    requests = workbook.active
    requests.title = "요청목록"
    requests["A1"] = "Template"
    requests["A2"] = "1행 = 폼 1건. 실행=Y인 행만 처리."
    for column, name in {
        1: "실행(Y/N)", 2: "요청ID", 6: "폼 모델명(정확한 표기)",
        8: "신청 총 TPM(기존+추가)", 10: "구독 ID(행별 지정 시)",
    }.items():
        requests.cell(3, column, name)
    requests["A4"] = "Y"
    requests["B4"] = "SYNTHETIC-001"
    requests["F4"] = "Example-Model"
    requests["H4"] = 1000000
    requests["I4"] = '=IF(H4="","",H4/1000)'
    requests["L4"] = '="NA"'
    validation = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    requests.add_data_validation(validation)
    validation.add(requests["A4"])
    common = workbook.create_sheet("공통정보")
    common.append(["항목", "값", "설명"])
    common.append(["First Name", "FictionalFirst", "label"])
    common.append(["Last Name", "FictionalLast", "label"])
    common.append(["Company Email", "fictional@example.invalid", "label"])
    common["B4"].hyperlink = "mailto:fictional@example.invalid"
    common["B4"].comment = Comment("Private synthetic comment", "Synthetic Person")
    common.append(["Company Name", "Fictional Customer", "label"])
    reference = workbook.create_sheet("입력 포맷")
    reference["A1"] = "Example: Fictional Customer"
    workbook.create_sheet("폼 선택값")["A1"] = "Public reference"
    workbook.properties.creator = "Synthetic Person"
    workbook.save(path)
    workbook.close()


def test_blank_template_removes_values_hyperlinks_comments_and_metadata(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "template.xlsx"
    synthetic_source(source)
    original_bytes = source.read_bytes()
    result = make_template(source, output)
    assert result["status"] == "success"
    assert source.read_bytes() == original_bytes
    workbook = load_workbook(output)
    assert workbook["공통정보"]["B4"].value is None
    assert workbook["공통정보"]["B4"].hyperlink is None
    assert workbook["요청목록"]["I4"].value == '=IF(H4="","",H4/1000)'
    assert workbook["요청목록"]["L4"].value == '="NA"'
    assert "승인" in workbook["요청목록"]["A2"].value
    assert workbook["입력 포맷"]["A1"].value == "Example: Example Corporation"
    assert all(cell.comment is None for sheet in workbook for row in sheet for cell in row)
    assert workbook.properties.creator == "Foundry Quota Increase Skill"
    workbook.close()
    with ZipFile(output) as archive:
        all_xml = b"\n".join(
            archive.read(name) for name in archive.namelist()
            if name.endswith((".xml", ".rels"))
        )
        assert b"fictional@example.invalid" not in all_xml
        assert b"Synthetic Person" not in all_xml
        assert b"Fictional Customer" not in all_xml


def test_original_file_is_not_overwritten(tmp_path):
    source = tmp_path / "source.xlsx"
    synthetic_source(source)
    with pytest.raises(ValueError, match="not be overwritten"):
        make_template(source, source)


def test_package_metadata_cleanup_preserves_all_other_parts(tmp_path):
    original = tmp_path / "original.xlsx"
    source = tmp_path / "with-path.xlsx"
    synthetic_source(original)
    with ZipFile(original) as archive, ZipFile(source, "w") as output:
        for entry in archive.infolist():
            data = archive.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                document = minidom.parseString(data)
                node = document.createElement("absPath")
                node.setAttribute("url", "https://example.invalid/personal/synthetic-user/")
                document.documentElement.appendChild(node)
                data = document.toxml(encoding="utf-8")
            output.writestr(entry, data)
    assert sanitize_package_metadata(source) == 1
    assert sanitize_package_metadata(source) == 0
    with ZipFile(original) as before, ZipFile(source) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "xl/workbook.xml":
                assert before.read(name) == after.read(name)
        assert b"synthetic-user" not in after.read("xl/workbook.xml")
    workbook = load_workbook(source)
    assert workbook["요청목록"]["I4"].data_type == "f"
    workbook.close()


def test_bundled_template_is_blank_but_references_and_cached_formulas_remain():
    path = ROOT / "assets" / "quota_increase_request_template.xlsx"
    with ZipFile(path) as archive:
        document = minidom.parseString(archive.read("xl/workbook.xml"))
        assert not document.getElementsByTagNameNS("*", "absPath")
    workbook = load_workbook(path)
    cached = load_workbook(path, data_only=True)
    assert workbook.sheetnames == ["요청목록", "공통정보", "입력 포맷", "폼 선택값"]
    request = workbook["요청목록"]
    for row in range(4, request.max_row + 1):
        for column in (1, 2, 3, 4, 5, 6, 7, 8, 10, 11):
            assert request.cell(row, column).value is None
        assert request.cell(row, 9).data_type == "f"
        assert all(request.cell(row, col).data_type == "f" for col in range(12, 26))
        assert cached["요청목록"].cell(row, 12).value == "NA"
        assert cached["요청목록"].cell(row, 22).value == "L_NA"
    assert all(
        workbook["공통정보"].cell(row, 2).value is None
        for row in range(2, workbook["공통정보"].max_row + 1)
    )
    assert len(request.data_validations.dataValidation) == 7
    assert len(workbook.defined_names) == 51
    assert workbook["폼 선택값"]["I4"].value
    assert workbook["입력 포맷"].max_row >= 30
    assert workbook.properties.creator == "Foundry Quota Increase Skill"
    assert workbook.properties.lastModifiedBy == "Foundry Quota Increase Skill"
    assert not [
        cell for sheet in cached for row in sheet for cell in row if cell.data_type == "e"
    ]
    workbook.close()
    cached.close()

"""Synthetic-only intake and ledger tests; no browser or network operations."""

import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import Workbook, load_workbook
import pytest


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "quota_workbook.py"
SPEC = importlib.util.spec_from_file_location("quota_workbook", SCRIPT)
quota = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(quota)
SUBSCRIPTION = "11111111-2222-3333-4444-555555555555"
OTHER_SUBSCRIPTION = "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"


def common():
    return {
        "First Name": "Synthetic",
        "Last Name": "Tester",
        "Company Email": "quota-test@example.invalid",
        "Company Name": "Synthetic Example",
        "Company Address": "1 Test Street",
        "Company City": "Example City",
        "Company Postal Code": "00123",
        "Company Country": "Test Country",
        "Subscription Id": SUBSCRIPTION,
        "Justification": "Synthetic test only; no request will be submitted.",
    }


def request(request_id="REQ-001", **changes):
    result = {
        "execute": "Y", "request_id": request_id,
        "quota_type": "Model Deployment (PTU/RPM/TPM)",
        "model_type": "Azure OpenAI", "deployment_type": "Global Standard",
        "model_name": "synthetic-chat-model", "region": "synthetic-region",
        "requested_tpm": 120000, "subscription_id": "",
    }
    result.update(changes)
    return result


def write_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return path


def intake(tmp_path, rows=None, mode="submit", values=None, **kwargs):
    source = write_json(tmp_path / ("intake-" + quota.uuid.uuid4().hex + ".json"),
                        {"common": common() if values is None else values,
                         "requests": [request()] if rows is None else rows})
    run = tmp_path / ("run-" + quota.uuid.uuid4().hex)
    quota.prepare_json(source, run, mode=mode, **kwargs)
    return run


def row(run, request_id="REQ-001"):
    return quota.find_row(quota.load_manifest(run), request_id)


def screenshot(tmp_path, name="synthetic.png"):
    path = tmp_path / name
    path.write_bytes(b"\x89PNG\r\n\x1a\nsynthetic fixture")
    return str(path.resolve())


def ui_evidence(tmp_path, run, target_id="REQ-001", **changes):
    planned = row(run, target_id)
    result = {
        "request_id": target_id, "payload_hash": planned["payload_hash"], "status": "filled",
        **{field: planned[field] for field in
           ("quota_type", "model_type", "deployment_type", "model_name", "quota_input")},
        "region_asked": True, "region_selected": planned["region"],
        "common_verified": True, "subscription_id_verified": True,
        "screenshot_paths": [screenshot(tmp_path)], "detail": "Synthetic fixture only.",
    }
    result.update(changes)
    return result


def fill(tmp_path, run, target_id="REQ-001", **changes):
    evidence = write_json(tmp_path / ("ui-" + quota.uuid.uuid4().hex + ".json"),
                          ui_evidence(tmp_path, run, target_id, **changes))
    return quota.record_ui(run, target_id, evidence)


def approval(run):
    shown = quota.preview(run)
    quota.approve(run, shown["preview_fingerprint"], "Human reviewed these exact synthetic rows.")
    return shown


def begun(tmp_path, run, request_id="REQ-001"):
    fill(tmp_path, run, request_id)
    approval(run)
    return quota.begin_submit(run, request_id)


def submission_evidence(tmp_path, run, target_id="REQ-001", **changes):
    planned = row(run, target_id)
    evidence = {
        "request_id": target_id, "attempt_id": planned["attempts"][-1]["attempt_id"],
        "payload_hash": planned["payload_hash"], "submission_status": "submitted",
        "submit_clicked": True, "confirmation_text": "Your response was submitted.",
        "confirmation_screenshot": screenshot(tmp_path, "confirmation.png"),
    }
    evidence.update(changes)
    return write_json(tmp_path / ("finish-" + quota.uuid.uuid4().hex + ".json"), evidence)


def rewrite_manifest(run, function):
    manifest = quota.load_manifest(run)
    function(manifest)
    quota.atomic_json(quota.manifest_path(run), manifest)


def workbook(tmp_path, rows=None, common_values=None, alias=False, header_order=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "요청 목록" if alias else "요청목록"
    sheet.append(["Synthetic test workbook"])
    sheet.append([])
    fields = list(quota.INPUT_FIELDS if header_order is None else header_order)
    header_by_field = {field: header for header, field in quota.WORKBOOK_HEADERS.items()}
    sheet.append([header_by_field[field] for field in fields])
    for item in [request()] if rows is None else rows:
        sheet.append([item.get(field) for field in fields])
    for index in range(10, 15):
        sheet.cell(index, fields.index("execute") + 1, "N")
        sheet.cell(index, fields.index("quota_input") + 1, "=IF(H10=\"\",\"\",H10/1000)")
        sheet.cell(index, 12, "=1+1")
    sheet.cell(20, fields.index("quota_input") + 1, "=1+1")
    values = book.create_sheet("공통 정보" if alias else "공통정보")
    values.append(["항목", "값", "설명"])
    for label, value in (common() if common_values is None else common_values).items():
        values.append([label, value])
    reference = book.create_sheet("입력포맷" if alias else "입력 포맷")
    reference.append(["Reference only"])
    reference.cell(4, 1, "Y")
    reference.cell(4, 2, "MUST-NOT-READ")
    book.create_sheet("폼 선택값").append(["Reference only; no authoritative saved choices."])
    path = tmp_path / ("workbook-" + quota.uuid.uuid4().hex + ".xlsx")
    book.save(path)
    book.close()
    return path


def test_workbook_readonly_aliases_units_formulas_and_template_rows(tmp_path):
    source = workbook(tmp_path, rows=[request(quota_input="=999999")], alias=True)
    original = source.read_bytes()
    run = tmp_path / "run"
    result = quota.prepare(source, run)
    assert source.read_bytes() == original
    assert result["mode"] == "capture"
    assert result["approval"] is None
    assert result["source"]["sha256"] == quota.hashlib.sha256(original).hexdigest()
    assert result["created_at"].endswith("Z")
    assert len(result["requests"]) == 1
    planned = result["requests"][0]
    assert planned["source_row"] == 4
    assert planned["requested_tpm"] == 120000
    assert planned["quota_input"] == "120"
    assert planned["status"] == "pending"
    assert result["common"]["Company Postal Code"] == "00123"
    book = load_workbook(source, read_only=True, data_only=False)
    assert "공통 정보" in book.sheetnames and "입력포맷" in book.sheetnames
    book.close()


def test_reference_sheet_cannot_replace_missing_requests_sheet(tmp_path):
    source = workbook(tmp_path)
    book = load_workbook(source)
    del book["요청목록"]
    book.save(source)
    book.close()
    with pytest.raises(quota.LedgerError, match="요청목록"):
        quota.prepare(source, tmp_path / "run")


def test_workbook_reordered_columns_use_exact_header_names_not_offsets(tmp_path):
    source = workbook(
        tmp_path, [request(requested_tpm=230000, subscription_id=OTHER_SUBSCRIPTION,
                           notes="Synthetic reordered input.")],
        header_order=list(reversed(quota.INPUT_FIELDS)))
    original = source.read_bytes()
    result = quota.prepare(source, tmp_path / "run")
    planned = result["requests"][0]
    assert len(result["requests"]) == 1
    assert planned["status"] == "pending"
    assert planned["model_name"] == "synthetic-chat-model"
    assert planned["subscription_id"] == OTHER_SUBSCRIPTION
    assert planned["requested_tpm"] == 230000
    assert planned["quota_input"] == "230"
    assert planned["notes"] == "Synthetic reordered input."
    assert source.read_bytes() == original


@pytest.mark.parametrize("duplicate", [False, True])
def test_workbook_missing_or_duplicate_exact_headers_rejected(tmp_path, duplicate):
    source = workbook(tmp_path)
    book = load_workbook(source)
    book["요청목록"].cell(3, 11, "폼 모델명(정확한 표기)" if duplicate else "Renamed notes header")
    book.save(source)
    book.close()
    with pytest.raises(quota.LedgerError, match="header"):
        quota.prepare(source, tmp_path / "run")


def test_reordered_selected_input_formulas_still_rejected(tmp_path):
    source = workbook(tmp_path, [request(model_name="='[external.xlsx]Sheet1'!A1")],
                      header_order=list(reversed(quota.INPUT_FIELDS)))
    result = quota.prepare(source, tmp_path / "run")
    assert result["requests"][0]["status"] == "invalid"
    assert any("formula" in issue for issue in result["requests"][0]["issues"])


def test_whitespace_alias_ambiguity_rejected(tmp_path):
    source = workbook(tmp_path)
    book = load_workbook(source)
    book.create_sheet("공통 정보")
    book.save(source)
    book.close()
    with pytest.raises(quota.LedgerError, match="ambiguous"):
        quota.prepare(source, tmp_path / "run")


@pytest.mark.parametrize("value", [True, False, 0, -1000, 1001, 999, 1.5, 1000.0,
                                  "1,000", "1000.0", "1e3", "", None, "=1000",
                                  9007199254741000, "9007199254741000"])
def test_exact_tpm_rejects_unsupported_values(tmp_path, value):
    run = intake(tmp_path, [request(requested_tpm=value)])
    assert row(run)["status"] == "invalid"
    assert any("requested_tpm" in issue for issue in row(run)["issues"])


@pytest.mark.parametrize("value,expected", [
    (1000, "1"), (" 120000 ", "120"), (1000000000000000, "1000000000000"),
    (9007199254740000, "9007199254740"),
])
def test_exact_tpm_normalization(tmp_path, value, expected):
    run = intake(tmp_path, [request(requested_tpm=value)])
    assert row(run)["status"] == "pending"
    assert row(run)["quota_input"] == expected


@pytest.mark.parametrize("field", ["request_id", "quota_type", "model_type", "deployment_type",
                                   "model_name", "region", "subscription_id", "notes"])
def test_formulas_in_workbook_inputs_are_rejected(tmp_path, field):
    source = workbook(tmp_path, [request(**{field: "='[external.xlsx]Sheet1'!A1"})])
    result = quota.prepare(source, tmp_path / "run")
    assert result["requests"][0]["status"] == "invalid"
    assert any("formula" in issue for issue in result["requests"][0]["issues"])


def test_common_formula_is_not_evaluated(tmp_path):
    values = common()
    values["Company Name"] = "=HYPERLINK(\"https://example.invalid\",\"synthetic\")"
    source = workbook(tmp_path, common_values=values)
    result = quota.prepare(source, tmp_path / "run")
    assert result["requests"][0]["status"] == "invalid"
    assert result["common"]["Company Name"] == ""


@pytest.mark.parametrize("field,value", [("Company Email", "not-an-email"),
                                        ("First Name", ""), ("Company Postal Code", None),
                                        ("Subscription Id", "not-a-guid"),
                                        ("Subscription Id", quota.ZERO_GUID),
                                        ("Company City", True)])
def test_common_required_email_guid_and_boolean(tmp_path, field, value):
    values = common()
    values[field] = value
    run = intake(tmp_path, values=values)
    assert row(run)["status"] == "invalid"


def test_subscription_overrides_allow_blank_common(tmp_path):
    values = common()
    values["Subscription Id"] = ""
    run = intake(tmp_path, [request(subscription_id=OTHER_SUBSCRIPTION)], values=values)
    assert row(run)["status"] == "pending"
    assert row(run)["subscription_id"] == OTHER_SUBSCRIPTION
    run = intake(tmp_path, [request(subscription_id=OTHER_SUBSCRIPTION), request("REQ-002")], values=values)
    assert row(run)["status"] == "pending"
    assert row(run, "REQ-002")["status"] == "invalid"


def test_zero_guid_row_override_is_invalid_but_valid_override_replaces_placeholder_common(tmp_path):
    run = intake(tmp_path, [request(subscription_id=quota.ZERO_GUID)])
    assert row(run)["status"] == "invalid"
    assert any("nonzero GUID" in issue for issue in row(run)["issues"])
    values = common()
    values["Subscription Id"] = quota.ZERO_GUID
    run = intake(tmp_path, [request(subscription_id=OTHER_SUBSCRIPTION)], values=values)
    assert row(run)["status"] == "pending"
    assert row(run)["subscription_id"] == OTHER_SUBSCRIPTION


@pytest.mark.parametrize("model", ["synthetic-audio", "synthetic-realtime", "synthetic-real-time",
                                  "synthetic-live", "whisper", "synthetic-tts",
                                  "synthetic-transcribe", "synthetic-translate",
                                  "text-embedding-synthetic", "dall-e-synthetic",
                                  "synthetic-image", "synthetic-video", "sora-synthetic"])
def test_unsupported_openai_model_branches(tmp_path, model):
    run = intake(tmp_path, [request(model_name=model)])
    assert row(run)["status"] == "invalid"


@pytest.mark.parametrize("changes", [
    {"quota_type": "Tokens per Minute"}, {"quota_type": "PTU"},
    {"deployment_type": "Global Provisioned"}, {"deployment_type": "RPM"},
    {"deployment_type": "Standard"}, {"deployment_type": "Data Zone Standard"},
    {"model_type": "Other"}, {"model_type": "Azure Direct Model (ADM)", "deployment_type": "Standard"},
])
def test_unsupported_type_deployment_and_unit(tmp_path, changes):
    run = intake(tmp_path, [request(**changes)])
    assert row(run)["status"] == "invalid"


@pytest.mark.parametrize("changes", [
    {"deployment_type": "Global Standard"},
    {"model_type": "Azure Direct Model (ADM)", "model_name": "synthetic-adm-model"},
])
def test_supported_chat_tpm_and_adm_global_standard(tmp_path, changes):
    run = intake(tmp_path, [request(**changes)])
    assert row(run)["status"] == "pending"


def test_missing_and_duplicate_ids_and_duplicate_payload(tmp_path):
    source = workbook(tmp_path, [request(request_id="")])
    result = quota.prepare(source, tmp_path / "workbook-run")
    assert result["requests"][0]["status"] == "invalid"
    run = intake(tmp_path, [request(), request(region="other-region")])
    assert all(item["status"] == "invalid" for item in quota.load_manifest(run)["requests"])
    run = intake(tmp_path, [request(), request("REQ-002")])
    assert all("duplicate" in " ".join(item["issues"])
               for item in quota.load_manifest(run)["requests"])
    missing_id = request()
    del missing_id["request_id"]
    run = intake(tmp_path, [missing_id])
    assert row(run)["status"] == "pending"
    assert row(run)["source_row"] is None


@pytest.mark.parametrize("mode,include_disabled,ids,expected", [
    ("capture", False, None, "skipped"), ("capture", True, None, "pending"),
    ("capture", False, ["REQ-001"], "pending"), ("submit", False, None, "skipped"),
    ("submit", True, None, "skipped"), ("submit", True, ["REQ-001"], "skipped"),
])
def test_disabled_selection_never_becomes_submit_eligibility(tmp_path, mode, include_disabled, ids, expected):
    run = intake(tmp_path, [request(execute="N")], mode=mode,
                 include_disabled=include_disabled, ids=ids)
    assert row(run)["status"] == expected
    assert quota.preview(run)["ready_request_ids"] == []
    with pytest.raises(quota.LedgerError):
        quota.begin_submit(run, "REQ-001")


def test_ids_filter_and_unknown_id(tmp_path):
    rows = [request(), request("REQ-002", region="other-region")]
    run = intake(tmp_path, rows, ids=["REQ-002"])
    assert row(run)["status"] == "skipped"
    assert row(run, "REQ-002")["status"] == "pending"
    with pytest.raises(quota.LedgerError, match="do not exist"):
        intake(tmp_path, rows, ids=["REQ-999"])


def test_new_run_directory_required_and_ledger_preserved(tmp_path):
    source = workbook(tmp_path)
    run = tmp_path / "run"
    quota.prepare(source, run)
    original = quota.manifest_path(run).read_bytes()
    with pytest.raises(quota.LedgerError, match="already exists"):
        quota.prepare(source, run)
    assert quota.manifest_path(run).read_bytes() == original


def test_region_can_be_unasked_but_not_missing_when_live_form_asks(tmp_path):
    run = intake(tmp_path, [request(region=None)])
    assert row(run)["status"] == "pending"
    result = fill(tmp_path, run, region_asked=False, region_selected=None)
    assert result["status"] == "filled"
    assert quota.preview(run)["ready_request_ids"] == ["REQ-001"]
    result = fill(tmp_path, run, region_asked=True, region_selected=None)
    assert result["status"] == "blocked"


def test_preview_distinguishes_planned_region_from_effective_not_asked_region(tmp_path):
    run = intake(tmp_path)
    fill(tmp_path, run, region_asked=False, region_selected=None)
    shown = quota.preview(run)
    displayed = shown["requests"][0]
    assert displayed["classification"] == "ready"
    assert displayed["region"] == "synthetic-region"
    assert displayed["region_asked"] is False
    assert displayed["region_selected"] is None
    assert displayed["effective_region"] == "not asked"
    fill(tmp_path, run)
    asked = quota.preview(run)["requests"][0]
    assert asked["region_asked"] is True
    assert asked["region_selected"] == asked["effective_region"] == "synthetic-region"


def test_preview_marks_region_ui_facts_unverified_before_ui_evidence(tmp_path):
    run = intake(tmp_path)
    displayed = quota.preview(run)["requests"][0]
    assert displayed["region"] == "synthetic-region"
    assert displayed["region_asked"] is None
    assert displayed["region_selected"] is None
    assert displayed["effective_region"] == "unverified"


@pytest.mark.parametrize("initially_asked", [False, True])
def test_refill_changing_region_asked_invalidates_existing_approval(tmp_path, initially_asked):
    run = intake(tmp_path)
    fill(tmp_path, run, region_asked=initially_asked,
         region_selected="synthetic-region" if initially_asked else None)
    shown = approval(run)
    original_payload_hash = row(run)["payload_hash"]
    fill(tmp_path, run, region_asked=not initially_asked,
         region_selected=None if initially_asked else "synthetic-region")
    assert row(run)["status"] == "filled"
    assert row(run)["payload_hash"] == original_payload_hash
    assert quota.scope_fingerprint(quota.load_manifest(run), ["REQ-001"]) != shown["preview_fingerprint"]
    with pytest.raises(quota.LedgerError, match="stale"):
        quota.begin_submit(run, "REQ-001")


def test_effective_selected_region_is_bound_to_fingerprint(tmp_path):
    run = intake(tmp_path)
    fill(tmp_path, run)
    shown = approval(run)
    rewrite_manifest(
        run, lambda manifest: manifest["requests"][0]["ui_evidence"].update(
            {"region_selected": "different-synthetic-region"}))
    assert quota.scope_fingerprint(quota.load_manifest(run), ["REQ-001"]) != shown["preview_fingerprint"]
    with pytest.raises(quota.LedgerError, match="stale"):
        quota.begin_submit(run, "REQ-001")


def test_refill_with_same_effective_region_keeps_approval_valid(tmp_path):
    run = intake(tmp_path)
    fill(tmp_path, run, region_asked=False, region_selected=None)
    shown = approval(run)
    fill(tmp_path, run, region_asked=False, region_selected=None,
         screenshot_paths=[screenshot(tmp_path, "new-synthetic-capture.png")],
         detail="New synthetic screenshot and detail; identical effective submitted fields.",
         timestamp="2099-01-01T00:00:00Z")
    manifest = quota.load_manifest(run)
    assert quota.scope_fingerprint(manifest, ["REQ-001"]) == shown["preview_fingerprint"]
    assert quota.begin_submit(run, "REQ-001")["status"] == "submitting"


def test_submission_status_and_evidence_metadata_do_not_change_region_bound_scope(tmp_path):
    run = intake(tmp_path, [request(), request("REQ-002", region="other-region")])
    fill(tmp_path, run, region_asked=False, region_selected=None)
    fill(tmp_path, run, "REQ-002")
    shown = approval(run)
    quota.begin_submit(run, "REQ-001")
    quota.finish_submit(run, "REQ-001", submission_evidence(tmp_path, run))
    manifest = quota.load_manifest(run)
    assert quota.scope_fingerprint(manifest, ["REQ-001", "REQ-002"]) == shown["preview_fingerprint"]
    assert quota.begin_submit(run, "REQ-002")["status"] == "submitting"


@pytest.mark.parametrize("changes", [
    {"payload_hash": "wrong"}, {"request_id": "wrong"}, {"model_name": "wrong"},
    {"quota_type": "wrong"}, {"model_type": "wrong"}, {"deployment_type": "wrong"},
    {"quota_input": "120.0"}, {"quota_input": 120},
    {"region_selected": "wrong"}, {"region_asked": "true"},
    {"region_asked": False, "region_selected": "synthetic-region"},
    {"common_verified": False}, {"common_verified": 1},
    {"subscription_id_verified": False}, {"subscription_id": OTHER_SUBSCRIPTION},
    {"screenshot_paths": []}, {"screenshot_paths": ["relative.png"]},
    {"screenshot_paths": ["C:\\nonexistent-synthetic-image.png"]},
    {"visible_errors": ["Synthetic form error"]},
    {"errors": ["Synthetic error"]}, {"validation_errors": ["Synthetic error"]},
    {"status": "model_missing"}, {"status": "region_missing"},
    {"status": "validation_error"}, {"status": "unsupported_form"},
])
def test_bad_ui_evidence_never_readies_row(tmp_path, changes):
    run = intake(tmp_path)
    result = fill(tmp_path, run, **changes)
    assert result["status"] == "blocked"
    assert quota.preview(run)["ready_request_ids"] == []


def test_capture_can_fill_but_never_approve(tmp_path):
    run = intake(tmp_path, mode="capture")
    fill(tmp_path, run)
    shown = quota.preview(run)
    assert shown["ready_request_ids"] == ["REQ-001"]
    with pytest.raises(quota.LedgerError, match="Capture"):
        quota.approve(run, shown["preview_fingerprint"], "Human reviewed.")
    with pytest.raises(quota.LedgerError):
        quota.begin_submit(run, "REQ-001")


def test_explicit_preview_and_nonempty_confirmation_required(tmp_path):
    run = intake(tmp_path)
    fill(tmp_path, run)
    fingerprint = quota.scope_fingerprint(quota.load_manifest(run), ["REQ-001"])
    with pytest.raises(quota.LedgerError, match="Preview"):
        quota.approve(run, fingerprint, "Human reviewed.")
    quota.preview(run)
    with pytest.raises(quota.LedgerError, match="Nonblank"):
        quota.approve(run, fingerprint, " ")
    with pytest.raises(quota.LedgerError, match="incorrect"):
        quota.approve(run, "wrong", "Human reviewed.")
    with pytest.raises(quota.LedgerError, match="approval"):
        quota.begin_submit(run, "REQ-001")
    quota.approve(run, fingerprint, "Human reviewed.")
    with pytest.raises(quota.LedgerError, match="already has"):
        quota.approve(run, fingerprint, "Human reviewed again.")


def test_preview_displays_invalid_and_skipped_rows_outside_ready_scope(tmp_path):
    run = intake(tmp_path, [
        request(), request("REQ-002", requested_tpm=0),
        request("REQ-003", execute="N", region="other-region"),
    ])
    fill(tmp_path, run)
    shown = approval(run)
    assert shown["common"] == common()
    assert [item["classification"] for item in shown["requests"]] == ["ready", "blocked", "skipped"]
    assert quota.load_manifest(run)["approval"]["request_ids"] == ["REQ-001"]


def test_newly_filled_row_does_not_expand_approved_scope(tmp_path):
    run = intake(tmp_path, [request(), request("REQ-002", region="other-region")])
    fill(tmp_path, run)
    approval(run)
    fill(tmp_path, run, "REQ-002")
    with pytest.raises(quota.LedgerError, match="approval"):
        quota.begin_submit(run, "REQ-002")
    assert quota.begin_submit(run, "REQ-001")["status"] == "submitting"


def test_preview_becomes_stale_if_ready_scope_changes_before_approval(tmp_path):
    run = intake(tmp_path, [request(), request("REQ-002", region="other-region")])
    fill(tmp_path, run)
    shown = quota.preview(run)
    fill(tmp_path, run, "REQ-002")
    with pytest.raises(quota.LedgerError, match="Preview"):
        quota.approve(run, shown["preview_fingerprint"], "Human reviewed only first row.")


@pytest.mark.parametrize("target", ["model_name", "requested_tpm", "common", "run_id", "scope"])
def test_immutable_tampering_invalidates_approval(tmp_path, target):
    run = intake(tmp_path)
    fill(tmp_path, run)
    approval(run)

    def tamper(manifest):
        if target == "common":
            manifest["common"]["Justification"] = "Changed synthetic justification."
        elif target == "scope":
            manifest["approval"]["request_ids"].append("REQ-999")
        elif target == "run_id":
            manifest["run_id"] = "other-run"
        else:
            planned = manifest["requests"][0]
            planned[target] = "changed-model" if target == "model_name" else 130000
            # Rehashing an edited row cannot refresh the old human approval.
            planned["payload_hash"] = quota.payload_hash(manifest["common"], planned)
    rewrite_manifest(run, tamper)
    with pytest.raises(quota.LedgerError):
        quota.begin_submit(run, "REQ-001")


def test_status_changes_do_not_invalidate_remaining_batch_approval(tmp_path):
    run = intake(tmp_path, [request(), request("REQ-002", region="other-region")])
    fill(tmp_path, run)
    fill(tmp_path, run, "REQ-002")
    shown = approval(run)
    first = quota.begin_submit(run, "REQ-001")
    ticket = quota.read_json(first["ticket_path"])
    assert ticket["approved"] is True and ticket["common"] == common()
    assert ticket["approval_fingerprint"] == shown["preview_fingerprint"]
    assert quota.finish_submit(run, "REQ-001", submission_evidence(tmp_path, run))["status"] == "submitted"
    assert quota.begin_submit(run, "REQ-002")["status"] == "submitting"


def test_active_unknown_and_submitted_states_block_replay(tmp_path):
    run = intake(tmp_path)
    begun(tmp_path, run)
    with pytest.raises(quota.LedgerError):
        quota.begin_submit(run, "REQ-001")
    with pytest.raises(quota.LedgerError, match="cannot be overwritten"):
        fill(tmp_path, run)
    quota.mark_unknown(run, "REQ-001", "Synthetic browser timeout after click boundary.")
    assert quota.status(run)["stop_and_reconcile"] is True
    with pytest.raises(quota.LedgerError):
        quota.begin_submit(run, "REQ-001")
    with pytest.raises(quota.LedgerError):
        quota.finish_submit(run, "REQ-001", submission_evidence(tmp_path, run))
    with pytest.raises(quota.LedgerError):
        fill(tmp_path, run)


@pytest.mark.parametrize("confirmation", [
    "Your response was submitted.", "Your response has been submitted",
    "응답이 제출되었습니다.", "응답을 제출했습니다",
    "Thanks! You've completed the request!",
])
def test_successful_finish_is_not_replayable(tmp_path, confirmation):
    run = intake(tmp_path)
    begun(tmp_path, run)
    evidence = submission_evidence(tmp_path, run, confirmation_text=confirmation)
    assert quota.finish_submit(run, "REQ-001", evidence)["status"] == "submitted"
    with pytest.raises(quota.LedgerError):
        quota.finish_submit(run, "REQ-001", evidence)
    with pytest.raises(quota.LedgerError):
        quota.begin_submit(run, "REQ-001")
    with pytest.raises(quota.LedgerError):
        fill(tmp_path, run)


@pytest.mark.parametrize("changes", [
    {"confirmation_text": "Microsoft Foundry Service: Request for Quota Increase"},
    {"confirmation_text": "Thank you"}, {"confirmation_text": "Not Your response was submitted"},
    {"confirmation_text": "Thanks!"},
    {"confirmation_text": "Not Thanks! You've completed the request!"},
    {"confirmation_text": ""}, {"confirmation_screenshot": None},
    {"confirmation_screenshot": "relative.png"}, {"submit_clicked": False},
    {"submit_clicked": 1}, {"submission_status": "unknown"},
    {"submission_status": "not_clicked", "submit_clicked": True, "detail": "Something failed."},
    {"visible_errors": ["Synthetic error after click"]},
])
def test_ambiguous_or_failed_click_finishes_unknown(tmp_path, changes):
    run = intake(tmp_path)
    begun(tmp_path, run)
    result = quota.finish_submit(run, "REQ-001", submission_evidence(tmp_path, run, **changes))
    assert result["status"] == "unknown"
    assert result["stop_and_reconcile"] is True


@pytest.mark.parametrize("changes", [{"request_id": "wrong"}, {"attempt_id": "wrong"}, {"payload_hash": "wrong"}])
def test_submission_evidence_must_match_attempt(tmp_path, changes):
    run = intake(tmp_path)
    begun(tmp_path, run)
    with pytest.raises(quota.LedgerError, match="does not match"):
        quota.finish_submit(run, "REQ-001", submission_evidence(tmp_path, run, **changes))
    assert row(run)["status"] == "submitting"


def test_explicit_not_clicked_allows_manual_new_attempt_not_old_ticket(tmp_path):
    run = intake(tmp_path)
    first = begun(tmp_path, run)
    evidence = submission_evidence(tmp_path, run, submission_status="not_clicked", submit_clicked=False,
                                   detail="Synthetic harness confirms no click was dispatched.")
    assert quota.finish_submit(run, "REQ-001", evidence)["status"] == "filled"
    assert row(run)["attempts"][-1]["status"] == "not_clicked"
    second = quota.begin_submit(run, "REQ-001")
    assert first["attempt_id"] != second["attempt_id"]
    assert len(row(run)["attempts"]) == 2


def test_not_clicked_without_explicit_detail_is_unknown(tmp_path):
    run = intake(tmp_path)
    begun(tmp_path, run)
    evidence = submission_evidence(tmp_path, run, submission_status="not_clicked", submit_clicked=False)
    assert quota.finish_submit(run, "REQ-001", evidence)["status"] == "unknown"


def test_global_unknown_guard_blocks_other_rows(tmp_path):
    run = intake(tmp_path, [request(), request("REQ-002", region="other-region")])
    fill(tmp_path, run)
    fill(tmp_path, run, "REQ-002")
    approval(run)
    quota.begin_submit(run, "REQ-001")
    quota.mark_unknown(run, "REQ-001", "Synthetic timeout.")
    with pytest.raises(quota.LedgerError, match="reconciliation"):
        quota.begin_submit(run, "REQ-002")


def unknown_rendered(tmp_path):
    run = intake(tmp_path)
    started = begun(tmp_path, run)
    quota.render(run, "REQ-001", "submit", tmp_path / "unknown-submit.js",
                 tmp_path, started["ticket_path"])
    quota.mark_unknown(run, "REQ-001", "Synthetic interrupted browser call after click boundary.")
    return run


def test_reconcile_submitted_preserves_unknown_evidence_and_never_replays(tmp_path):
    run = unknown_rendered(tmp_path)
    previous = copy.deepcopy(row(run)["attempts"][-1])
    evidence = submission_evidence(
        tmp_path, run, confirmation_text="Thanks! You've completed the request!")
    result = quota.reconcile_submitted(
        run, "REQ-001", evidence, "Human verified this exact synthetic completion screen.")
    assert result == {
        "request_id": "REQ-001", "status": "submitted",
        "new_clicks": 0, "stop_and_reconcile": False,
    }
    planned = row(run)
    assert len(planned["attempts"]) == 1
    assert planned["attempts"][0]["attempt_id"] == previous["attempt_id"]
    assert planned["attempts"][0]["detail"] == previous["detail"]
    assert planned["attempts"][0]["reconciliation"]["evidence"] == quota.read_json(evidence)
    with pytest.raises(quota.LedgerError):
        quota.begin_submit(run, "REQ-001")
    with pytest.raises(quota.LedgerError):
        quota.reconcile_submitted(run, "REQ-001", evidence, "Repeated confirmation.")


@pytest.mark.parametrize("changes", [
    {"request_id": "wrong"},
    {"attempt_id": "wrong"},
    {"payload_hash": "wrong"},
    {"confirmation_text": "Thanks!"},
    {"confirmation_screenshot": None},
    {"submit_clicked": False},
    {"submission_status": "unknown"},
    {"visible_errors": ["Synthetic error"]},
])
def test_reconcile_submitted_rejects_weak_or_mismatched_evidence(tmp_path, changes):
    run = unknown_rendered(tmp_path)
    evidence = submission_evidence(tmp_path, run, **changes)
    with pytest.raises(quota.LedgerError):
        quota.reconcile_submitted(run, "REQ-001", evidence, "Human confirmation.")
    assert row(run)["status"] == "unknown"


def test_reconcile_submitted_requires_human_confirmation(tmp_path):
    run = unknown_rendered(tmp_path)
    with pytest.raises(quota.LedgerError, match="human confirmation"):
        quota.reconcile_submitted(run, "REQ-001", submission_evidence(tmp_path, run), "")
    assert row(run)["status"] == "unknown"


def test_reconcile_submitted_rejects_prior_evidence_of_no_click(tmp_path):
    run = unknown_rendered(tmp_path)
    rewrite_manifest(run, lambda manifest: manifest["requests"][0]["attempts"][0].update(
        {"evidence": {"submit_clicked": False}}))
    with pytest.raises(quota.LedgerError, match="not clicked"):
        quota.reconcile_submitted(
            run, "REQ-001", submission_evidence(tmp_path, run), "Human confirmation.")


def test_reconcile_submitted_rechecks_immutable_approval(tmp_path):
    run = unknown_rendered(tmp_path)
    evidence = submission_evidence(tmp_path, run)
    rewrite_manifest(run, lambda manifest: manifest["common"].update(
        {"Company City": "Different city"}))
    with pytest.raises(quota.LedgerError):
        quota.reconcile_submitted(run, "REQ-001", evidence, "Human confirmation.")


def test_status_counts_requests_not_maximum_id(tmp_path):
    run = intake(tmp_path, [
        request("REQ-001"),
        request("REQ-034", model_name="different-synthetic-model"),
    ])
    status = quota.status(run)
    assert status["total_requests"] == status["selected_count"] == 2
    assert status["submitted_count"] == 0
    assert status["counts"] == {"pending": 2}


def test_lock_prevents_duplicate_mutation(tmp_path):
    run = intake(tmp_path)
    fill(tmp_path, run)
    approval(run)
    lock = run / ".ledger.lock"
    lock.write_text("synthetic concurrent owner", encoding="utf-8")
    with pytest.raises(quota.LedgerError, match="locked"):
        quota.begin_submit(run, "REQ-001")
    assert row(run)["status"] == "filled"
    lock.unlink()
    assert quota.begin_submit(run, "REQ-001")["status"] == "submitting"


def test_atomic_replace_retries_transient_windows_sharing_failure(tmp_path, monkeypatch):
    path = write_json(tmp_path / "atomic.json", {"value": "old"})
    replace = quota.os.replace
    calls = []

    def briefly_locked(source, destination):
        calls.append(destination)
        if len(calls) == 1:
            raise PermissionError("Synthetic Windows sharing conflict.")
        return replace(source, destination)
    monkeypatch.setattr(quota.os, "replace", briefly_locked)
    monkeypatch.setattr(quota.time, "sleep", lambda delay: None)
    quota.atomic_json(path, {"value": "new"})
    assert len(calls) == 2
    assert quota.read_json(path) == {"value": "new"}
    assert not list(tmp_path.glob("*.tmp"))


def test_atomic_replace_exhaustion_preserves_original(tmp_path, monkeypatch):
    path = write_json(tmp_path / "atomic.json", {"value": "old"})

    def always_locked(source, destination):
        raise PermissionError("Synthetic persistent sharing conflict.")
    monkeypatch.setattr(quota.os, "replace", always_locked)
    monkeypatch.setattr(quota.time, "sleep", lambda delay: None)
    with pytest.raises(PermissionError):
        quota.atomic_json(path, {"value": "new"})
    assert quota.read_json(path) == {"value": "old"}
    assert not list(tmp_path.glob("*.tmp"))


def test_existing_ticket_is_not_overwritten_and_failed_publication_stops_retries(tmp_path, monkeypatch):
    run = intake(tmp_path)
    fill(tmp_path, run)
    approval(run)
    path = tmp_path / "ticket.json"
    path.write_text('{"synthetic":"existing"}', encoding="utf-8")
    with pytest.raises(quota.LedgerError, match="destination"):
        quota.begin_submit(run, "REQ-001", path)
    assert row(run)["status"] == "filled"
    assert path.read_text(encoding="utf-8") == '{"synthetic":"existing"}'
    actual = quota.atomic_json

    def fail_ticket(path, data, exclusive=False):
        if "attempt_id" in data:
            raise OSError("Synthetic publication failure.")
        return actual(path, data, exclusive)
    monkeypatch.setattr(quota, "atomic_json", fail_ticket)
    with pytest.raises(OSError):
        quota.begin_submit(run, "REQ-001")
    assert row(run)["status"] == "submitting"
    assert not (run / ".ledger.lock").exists()


@pytest.fixture
def browser_template(tmp_path, monkeypatch):
    template = tmp_path / "synthetic-browser-template.js"
    template.write_text("(async (page) => { const config = __QUOTA_CONFIG__; return config; })",
                        encoding="utf-8")
    monkeypatch.setattr(quota, "BROWSER_TEMPLATE", template)
    return template


def test_render_is_safe_json_injection_with_quotes_unicode_and_no_execution(tmp_path, browser_template):
    values = common()
    values["Justification"] = "Synthetic 'quotes' \"double\" 한글 \\ newline\n\u2028\u2029 end"
    run = intake(tmp_path, values=values, mode="capture")
    output = tmp_path / "rendered-fill.js"
    result = quota.render(run, "REQ-001", "fill", output, tmp_path / "screenshots")
    script = output.read_text(encoding="utf-8")
    assert result["action"] == "fill" and "JSON.parse(" in script
    encoded = script.split("JSON.parse(", 1)[1].rsplit("); return", 1)[0]
    config = json.loads(json.loads(encoded))
    assert config["common"] == values
    assert config["ticket"] is None and config["mode"] == "capture"
    assert config["request"]["payload_hash"] == row(run)["payload_hash"]
    assert config["request"]["payload_hash"] == quota.payload_hash(
        config["common"], config["request"])
    assert Path(config["screenshot_dir"]).is_absolute()
    assert row(run)["status"] == "pending"
    assert (tmp_path / "screenshots").is_dir()
    assert not list((tmp_path / "screenshots").iterdir())
    with pytest.raises(quota.LedgerError, match="output must be new"):
        quota.render(run, "REQ-001", "fill", output, tmp_path / "screenshots")


def test_submit_render_requires_exact_ticket_and_only_renders_once(tmp_path, browser_template):
    run = intake(tmp_path)
    with pytest.raises(quota.LedgerError, match="ticket"):
        quota.render(run, "REQ-001", "submit", tmp_path / "no-ticket.js", tmp_path)
    started = begun(tmp_path, run)
    ticket = quota.read_json(started["ticket_path"])
    wrong = copy.deepcopy(ticket)
    wrong["request"]["quota_input"] = "999"
    fake = write_json(tmp_path / "fake-ticket.json", wrong)
    with pytest.raises(quota.LedgerError, match="ticket does not match"):
        quota.render(run, "REQ-001", "submit", tmp_path / "wrong.js", tmp_path, fake)
    assert not (tmp_path / "wrong.js").exists()
    quota.render(run, "REQ-001", "submit", tmp_path / "submit.js", tmp_path, started["ticket_path"])
    script = (tmp_path / "submit.js").read_text(encoding="utf-8")
    encoded = script.split("JSON.parse(", 1)[1].rsplit("); return", 1)[0]
    config = json.loads(json.loads(encoded))
    assert config["request"]["status"] == "submitting"
    assert config["request"]["payload_hash"] == config["ticket"]["payload_hash"] == row(run)["payload_hash"]
    assert config["ticket"]["attempt_id"] == row(run)["attempts"][-1]["attempt_id"]
    assert row(run)["attempts"][-1]["rendered_at"]
    with pytest.raises(quota.LedgerError, match="already produced"):
        quota.render(run, "REQ-001", "submit", tmp_path / "replay.js", tmp_path, started["ticket_path"])
    with pytest.raises(quota.LedgerError):
        quota.render(run, "REQ-001", "fill", tmp_path / "fill.js", tmp_path)


def test_submit_render_rechecks_approval_common_and_ticket(tmp_path, browser_template):
    run = intake(tmp_path)
    started = begun(tmp_path, run)
    rewrite_manifest(run, lambda manifest: manifest["common"].update({"Company City": "Changed city"}))
    with pytest.raises(quota.LedgerError, match="changed"):
        quota.render(run, "REQ-001", "submit", tmp_path / "submit.js", tmp_path, started["ticket_path"])


def test_submit_mode_fill_render_needs_no_approval(tmp_path, browser_template):
    run = intake(tmp_path, mode="submit")
    assert quota.load_manifest(run)["approval"] is None
    assert quota.render(run, "REQ-001", "fill", tmp_path / "fill.js", tmp_path)["action"] == "fill"
    assert row(run)["status"] == "pending"


def test_render_creates_new_nested_screenshot_directory(tmp_path, browser_template):
    run = intake(tmp_path, mode="capture")
    captures = tmp_path / "new" / "nested" / "captures"
    quota.render(run, "REQ-001", "fill", tmp_path / "fill.js", captures)
    assert captures.is_dir()
    assert not list(captures.iterdir())


@pytest.mark.parametrize("action", ["fill", "submit"])
def test_render_rejects_screenshot_file_without_consuming_submit_attempt(
        tmp_path, browser_template, action):
    run = intake(tmp_path)
    started = begun(tmp_path, run) if action == "submit" else None
    captures = tmp_path / "captures"
    captures.write_text("Synthetic existing file; must not overwrite.", encoding="utf-8")
    output = tmp_path / "must-not-render.js"
    with pytest.raises(quota.LedgerError, match="not a directory"):
        quota.render(run, "REQ-001", action, output, captures,
                     started["ticket_path"] if started else None)
    assert not output.exists()
    assert captures.read_text(encoding="utf-8") == "Synthetic existing file; must not overwrite."
    if started:
        assert row(run)["status"] == "submitting"
        assert "rendered_at" not in row(run)["attempts"][-1]


def test_render_does_not_create_screenshot_directory_before_payload_validation(
        tmp_path, browser_template):
    run = intake(tmp_path, [request(deployment_type="Standard")])
    captures = tmp_path / "must-not-create" / "captures"
    with pytest.raises(quota.LedgerError):
        quota.render(run, "REQ-001", "fill", tmp_path / "invalid.js", captures)
    assert not captures.exists()


@pytest.mark.parametrize("terminal", ["unknown", "submitted"])
def test_submit_render_rejects_terminal_attempt_even_with_original_ticket(
        tmp_path, browser_template, terminal):
    run = intake(tmp_path)
    started = begun(tmp_path, run)
    if terminal == "unknown":
        quota.mark_unknown(run, "REQ-001", "Synthetic timeout: reconcile instead of replay.")
    else:
        quota.finish_submit(run, "REQ-001", submission_evidence(tmp_path, run))
    output = tmp_path / "must-not-render.js"
    with pytest.raises(quota.LedgerError):
        quota.render(run, "REQ-001", "submit", output, tmp_path, started["ticket_path"])
    assert not output.exists()


def test_submit_render_rejects_ticket_from_previous_not_clicked_attempt(tmp_path, browser_template):
    run = intake(tmp_path)
    first = begun(tmp_path, run)
    quota.finish_submit(
        run, "REQ-001", submission_evidence(
            tmp_path, run, submission_status="not_clicked", submit_clicked=False,
            detail="Synthetic harness verified that no click was dispatched."))
    quota.begin_submit(run, "REQ-001")
    with pytest.raises(quota.LedgerError, match="ticket does not match"):
        quota.render(run, "REQ-001", "submit", tmp_path / "old-attempt.js", tmp_path, first["ticket_path"])


@pytest.mark.parametrize("template", ["(async p => 1)", "__QUOTA_CONFIG__ __QUOTA_CONFIG__"])
def test_render_template_token_exactly_once(tmp_path, browser_template, template):
    browser_template.write_text(template, encoding="utf-8")
    run = intake(tmp_path)
    with pytest.raises(quota.LedgerError, match="exactly once"):
        quota.render(run, "REQ-001", "fill", tmp_path / "render.js", tmp_path)


def test_status_is_non_private(tmp_path):
    run = intake(tmp_path)
    result = json.dumps(quota.status(run))
    assert "pending" in result
    for value in common().values():
        assert value not in result


def test_cli_default_capture_preview_and_nonzero_errors(tmp_path):
    source = write_json(tmp_path / "synthetic-input.json", {"common": common(), "requests": [request()]})
    run = tmp_path / "cli-run"
    command = [sys.executable, str(SCRIPT)]
    prepared = subprocess.run(command + ["prepare-json", "--input", str(source), "--run-dir", str(run)],
                              text=True, capture_output=True, check=False)
    assert prepared.returncode == 0, prepared.stderr
    assert json.loads(prepared.stdout)["mode"] == "capture"
    assert common()["Company Email"] not in prepared.stdout
    shown = subprocess.run(command + ["preview", "--run-dir", str(run)],
                           text=True, capture_output=True, check=False)
    assert shown.returncode == 0
    assert json.loads(shown.stdout)["common"] == common()
    failed = subprocess.run(command + ["begin-submit", "--run-dir", str(run), "--id", "REQ-001"],
                            text=True, capture_output=True, check=False)
    assert failed.returncode != 0 and "Traceback" not in failed.stderr
    assert common()["Company Email"] not in failed.stderr


def test_cli_rejects_malformed_manifest_without_traceback_or_private_payload(tmp_path):
    run = intake(tmp_path)
    rewrite_manifest(run, lambda manifest: manifest.pop("mode"))
    completed = subprocess.run([sys.executable, str(SCRIPT), "status", "--run-dir", str(run)],
                               text=True, capture_output=True, check=False)
    assert completed.returncode != 0 and "Traceback" not in completed.stderr
    assert common()["Company Email"] not in completed.stderr


def test_concurrent_cli_begin_creates_only_one_durable_attempt_and_ticket(tmp_path):
    run = intake(tmp_path)
    fill(tmp_path, run)
    approval(run)
    command = [sys.executable, str(SCRIPT), "begin-submit", "--run-dir", str(run), "--id", "REQ-001"]
    first = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    second = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    first.communicate(timeout=30)
    second.communicate(timeout=30)
    assert sorted([first.returncode, second.returncode]) == [0, 2]
    assert row(run)["status"] == "submitting"
    assert len(row(run)["attempts"]) == 1
    assert len(list(run.glob("ticket-*.json"))) == 1
